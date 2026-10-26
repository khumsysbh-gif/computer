#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ໂປຣແກຣມບັນທຶກຖານຂໍ້ມູນບຸກຄະລາກອນ (Streamlit)
ເພີ່ມ / ແກ້ໄຂ / ລຶບ / ຄົ້ນຫາ / ສົ່ງອອກ Excel  |  3 ພາສາ: ລາວ, English, ไทย

ຕິດຕັ້ງ:  pip install streamlit pillow openpyxl pandas
ແລ່ນ:     python -m streamlit run database_app.py
"""
import hashlib
import hmac
import os
import re
import sqlite3
from datetime import datetime
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageOps

DB_FILE = "data.db"
PBKDF2_ROUNDS = 200_000
MAX_LOGIN_ATTEMPTS = 3
MIN_PASSWORD_LEN = 6
PHOTO_SIZE = (300, 400)

# ---------------------------------------------------------------- ຄຳສັບ 3 ພາສາ
LANG = {
    "lo": {
        "title": "ໂປຣແກຣມບັນທຶກຖານຂໍ້ມູນບຸກຄະລາກອນ",
        "form": "ຂໍ້ມູນ", "photo": "ຮູບ 3×4",
        "pick_photo": "ເລືອກຮູບ", "remove_photo": "ລຶບຮູບ",
        "add": "➕ ເພີ່ມ", "update": "✏️ ບັນທຶກການແກ້ໄຂ", "delete": "🗑 ລຶບ",
        "clear": "ລ້າງຟອມ", "export": "📥 ສ້າງໄຟລ໌ Excel", "download": "⬇️ ດາວໂຫຼດ Excel",
        "change_pw": "🔑 ປ່ຽນລະຫັດຜ່ານ", "logout": "🚪 ອອກຈາກລະບົບ",
        "search": "ຄົ້ນຫາ", "search_ph": "ພິມຄຳຄົ້ນຫາ...", "show_all": "ສະແດງທັງໝົດ",
        "no_photo": "ຍັງບໍ່ມີຮູບ",
        "required_hint": "* ຈຳເປັນຕ້ອງໃສ່     ຮູບແບບວັນທີ: ວວ/ດດ/ປປປປ  (ເຊັ່ນ 25/10/2026)",
        "select_row_err": "ກະລຸນາເລືອກແຖວຈາກຕາຕະລາງກ່ອນ",
        "success": "ສຳເລັດ", "error": "ຜິດພາດ",
        "setup_title": "ຕັ້ງລະຫັດຜ່ານຄັ້ງທຳອິດ",
        "new_pw": f"ລະຫັດຜ່ານໃໝ່ (ຢ່າງໜ້ອຍ {MIN_PASSWORD_LEN} ຕົວ)",
        "confirm_pw": "ຢືນຢັນລະຫັດຜ່ານ", "save": "ບັນທຶກ",
        "login_title": "ເຂົ້າສູ່ລະບົບ", "password": "ລະຫັດຜ່ານ", "login": "ເຂົ້າສູ່ລະບົບ",
        "wrong_pw": "ລະຫັດຜ່ານບໍ່ຖືກ", "attempts": "ເຫຼືອໂອກາດ",
        "locked": "ໃສ່ລະຫັດຜິດເກີນກຳນົດ. ກະລຸນາໂຫຼດໜ້າໃໝ່ເພື່ອລອງອີກ.",
        "pw_short": "ລະຫັດຜ່ານສັ້ນເກີນໄປ", "pw_mismatch": "ລະຫັດຜ່ານບໍ່ກົງກັນ",
        "pw_updated": "ປ່ຽນລະຫັດຜ່ານແລ້ວ", "old_pw": "ລະຫັດຜ່ານປັດຈຸບັນ",
        "name_required": "ກະລຸນາໃສ່ຊື່", "bad_date": "ຮູບແບບວັນທີບໍ່ຖືກ (ວວ/ດດ/ປປປປ):",
        "added": "ເພີ່ມຂໍ້ມູນສຳເລັດ", "updated": "ແກ້ໄຂຂໍ້ມູນສຳເລັດ", "deleted": "ລຶບຂໍ້ມູນສຳເລັດ",
        "confirm_delete": "ຢືນຢັນການລຶບຂໍ້ມູນນີ້?", "yes": "ຢືນຢັນ", "cancel": "ຍົກເລີກ",
        "total": "ທັງໝົດ", "no_data": "ບໍ່ມີຂໍ້ມູນ", "editing": "ກຳລັງແກ້ໄຂ ID",
        "language": "ພາສາ", "upload_hint": "ອັບໂຫຼດຮູບ (jpg, png, bmp)",
    },
    "en": {
        "title": "Personnel Database Management System",
        "form": "Information", "photo": "Photo 3×4",
        "pick_photo": "Select Photo", "remove_photo": "Remove Photo",
        "add": "➕ Add", "update": "✏️ Save Update", "delete": "🗑 Delete",
        "clear": "Clear Form", "export": "📥 Build Excel", "download": "⬇️ Download Excel",
        "change_pw": "🔑 Change Password", "logout": "🚪 Logout",
        "search": "Search", "search_ph": "Type to search...", "show_all": "Show All",
        "no_photo": "No photo",
        "required_hint": "* Required     Date format: DD/MM/YYYY  (e.g., 25/10/2026)",
        "select_row_err": "Please select a row from the table first",
        "success": "Success", "error": "Error",
        "setup_title": "First-time password setup",
        "new_pw": f"New password (min {MIN_PASSWORD_LEN} chars)",
        "confirm_pw": "Confirm password", "save": "Save",
        "login_title": "Login", "password": "Password", "login": "Login",
        "wrong_pw": "Incorrect password", "attempts": "Attempts left",
        "locked": "Too many failed attempts. Reload the page to try again.",
        "pw_short": "Password too short", "pw_mismatch": "Passwords do not match",
        "pw_updated": "Password updated", "old_pw": "Current password",
        "name_required": "Please input Name", "bad_date": "Invalid date format (DD/MM/YYYY):",
        "added": "Added successfully", "updated": "Updated successfully", "deleted": "Deleted successfully",
        "confirm_delete": "Delete this record?", "yes": "Confirm", "cancel": "Cancel",
        "total": "Total", "no_data": "No data", "editing": "Editing ID",
        "language": "Language", "upload_hint": "Upload photo (jpg, png, bmp)",
    },
    "th": {
        "title": "โปรแกรมบันทึกฐานข้อมูลบุคลากร",
        "form": "ข้อมูล", "photo": "รูป 3×4",
        "pick_photo": "เลือกรูป", "remove_photo": "ลบรูป",
        "add": "➕ เพิ่ม", "update": "✏️ บันทึกการแก้ไข", "delete": "🗑 ลบ",
        "clear": "ล้างฟอร์ม", "export": "📥 สร้างไฟล์ Excel", "download": "⬇️ ดาวน์โหลด Excel",
        "change_pw": "🔑 เปลี่ยนรหัสผ่าน", "logout": "🚪 ออกจากระบบ",
        "search": "ค้นหา", "search_ph": "พิมพ์คำค้นหา...", "show_all": "แสดงทั้งหมด",
        "no_photo": "ยังไม่มีรูป",
        "required_hint": "* จำเป็นต้องใส่     รูปแบบวันที่: วว/ดด/ปปปป  (เช่น 25/10/2026)",
        "select_row_err": "กรุณาเลือกแถวจากตารางก่อน",
        "success": "สำเร็จ", "error": "ผิดพลาด",
        "setup_title": "ตั้งรหัสผ่านครั้งแรก",
        "new_pw": f"รหัสผ่านใหม่ (อย่างน้อย {MIN_PASSWORD_LEN} ตัว)",
        "confirm_pw": "ยืนยันรหัสผ่าน", "save": "บันทึก",
        "login_title": "เข้าสู่ระบบ", "password": "รหัสผ่าน", "login": "เข้าสู่ระบบ",
        "wrong_pw": "รหัสผ่านไม่ถูกต้อง", "attempts": "โอกาสที่เหลือ",
        "locked": "ใส่รหัสผิดเกินกำหนด กรุณารีโหลดหน้าเพื่อลองใหม่",
        "pw_short": "รหัสผ่านสั้นเกินไป", "pw_mismatch": "รหัสผ่านไม่ตรงกัน",
        "pw_updated": "เปลี่ยนรหัสผ่านแล้ว", "old_pw": "รหัสผ่านปัจจุบัน",
        "name_required": "กรุณาใส่ชื่อ", "bad_date": "รูปแบบวันที่ไม่ถูกต้อง (วว/ดด/ปปปป):",
        "added": "เพิ่มข้อมูลสำเร็จ", "updated": "แก้ไขข้อมูลสำเร็จ", "deleted": "ลบข้อมูลสำเร็จ",
        "confirm_delete": "ยืนยันการลบข้อมูลนี้?", "yes": "ยืนยัน", "cancel": "ยกเลิก",
        "total": "ทั้งหมด", "no_data": "ไม่มีข้อมูล", "editing": "กำลังแก้ไข ID",
        "language": "ภาษา", "upload_hint": "อัปโหลดรูป (jpg, png, bmp)",
    },
}
LANG_NAMES = {"lo": "ລາວ", "en": "English", "th": "ไทย"}

FIELDS = [
    ("military_id", {"lo": "ເລກນາຍທະຫານ", "en": "Military ID", "th": "เลขนายทหาร"}, "text"),
    ("id_card", {"lo": "ເລກບັດປະຈຳຕົວ", "en": "ID Card", "th": "เลขบัตรประจำตัว"}, "text"),
    ("name", {"lo": "ຊື່", "en": "Name", "th": "ชื่อ"}, "text"),
    ("surname", {"lo": "ນາມສະກຸນ", "en": "Surname", "th": "นามสกุล"}, "text"),
    ("gender", {"lo": "ເພດ", "en": "Gender", "th": "เพศ"}, "combo"),
    ("rank", {"lo": "ຊັ້ນ", "en": "Rank", "th": "ชั้นยศ"}, "text"),
    ("birth_date", {"lo": "ວັນເດືອນປີເກີດ", "en": "Birth Date", "th": "วันเดือนปีเกิด"}, "date"),
    ("revolution_date", {"lo": "ວັນເດືອນປີເຂົ້າການປະຕິວັດ", "en": "Revolution Date", "th": "วันเดือนปีเข้าปฏิวัติ"}, "date"),
    ("promotion_date", {"lo": "ວັນເດືອນປີເລື່ອນຊັ້ນ", "en": "Promotion Date", "th": "วันเดือนปีเลื่อนชั้น"}, "date"),
    ("duty", {"lo": "ໜ້າທີ່ຮັບຜິດຊອບ", "en": "Duty", "th": "หน้าที่รับผิดชอบ"}, "text"),
    ("party_position", {"lo": "ຕຳແໜ່ງພັກ", "en": "Party Position", "th": "ตำแหน่งพรรค"}, "text"),
    ("state_position", {"lo": "ຕຳແໜ່ງລັດ", "en": "State Position", "th": "ตำแหน่งรัฐ"}, "text"),
    ("unit", {"lo": "ກົມກອງສັງກັດ", "en": "Unit", "th": "สังกัด/กรมกอง"}, "text"),
    ("reserve_party_date", {"lo": "ວັນເດືອນປີເຂົ້າພັກສຳຮອງ", "en": "Reserve Party Date", "th": "วันเข้าพรรคสำรอง"}, "date"),
    ("full_party_date", {"lo": "ວັນເດືອນປີເຂົ້າພັກສົມບູນ", "en": "Full Party Date", "th": "วันเข้าพรรคสมบูรณ์"}, "date"),
    ("education", {"lo": "ລະດັບການສຶກສາ", "en": "Education", "th": "ระดับการศึกษา"}, "combo"),
    ("phone", {"lo": "ເບີໂທ", "en": "Phone", "th": "เบอร์โทร"}, "text"),
    ("address", {"lo": "ທີ່ຢູ່", "en": "Address", "th": "ที่อยู่"}, "text"),
    ("note", {"lo": "ໝາຍເຫດ", "en": "Note", "th": "หมายเหตุ"}, "text"),
]
FIELD_KEYS = [f[0] for f in FIELDS]
DATE_KEYS = [f[0] for f in FIELDS if f[2] == "date"]

COMBO_OPTIONS = {
    "gender": {
        "lo": ["ຊາຍ", "ຍິງ"],
        "en": ["Male", "Female"],
        "th": ["ชาย", "หญิง"],
    },
    "education": {
        "lo": ["ປະຖົມ", "ມັດທະຍົມຕອນຕົ້ນ", "ມັດທະຍົມຕອນປາຍ", "ຊັ້ນກາງ / ວິຊາຊີບ", "ຊັ້ນສູງ", "ປະລິນຍາຕີ", "ປະລິນຍາໂທ", "ປະລິນຍາເອກ"],
        "en": ["Primary", "Lower Secondary", "Upper Secondary", "Vocational/Diploma", "Higher Diploma", "Bachelor's", "Master's", "Doctorate"],
        "th": ["ประถม", "มัธยมศึกษาตอนต้น", "มัธยมศึกษาตอนปลาย", "ปวส./อนุปริญญา", "ปริญญาตรี", "ปริญญาโท", "ปริญญาเอก"],
    },
}

TABLE_COLS = ["military_id", "name", "surname", "gender", "rank", "unit", "party_position", "phone"]


# ---------------------------------------------------------------- Database
class Database:
    def __init__(self, path=DB_FILE):
        # check_same_thread=False ເພາະ Streamlit ໃຊ້ຫຼາຍ thread
        self.conn = sqlite3.connect(path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self._create_tables()
        self._migrate()

    def _create_tables(self):
        with self.conn:
            self.conn.execute("""CREATE TABLE IF NOT EXISTS auth (
                id INTEGER PRIMARY KEY CHECK (id = 1), salt BLOB NOT NULL, pw_hash BLOB NOT NULL)""")
            self.conn.execute("""CREATE TABLE IF NOT EXISTS records (
                id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL)""")

    def _migrate(self):
        existing = {r["name"] for r in self.conn.execute("PRAGMA table_info(records)")}
        with self.conn:
            for key in FIELD_KEYS:
                if key not in existing:
                    self.conn.execute(f"ALTER TABLE records ADD COLUMN {key} TEXT")
            if "photo" not in existing:
                self.conn.execute("ALTER TABLE records ADD COLUMN photo BLOB")

    @staticmethod
    def _hash(password: str, salt: bytes) -> bytes:
        return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ROUNDS)

    def has_password(self) -> bool:
        return self.conn.execute("SELECT 1 FROM auth WHERE id = 1").fetchone() is not None

    def set_password(self, password: str):
        salt = os.urandom(16)
        with self.conn:
            self.conn.execute("INSERT OR REPLACE INTO auth (id, salt, pw_hash) VALUES (1, ?, ?)",
                              (salt, self._hash(password, salt)))

    def verify_password(self, password: str) -> bool:
        row = self.conn.execute("SELECT salt, pw_hash FROM auth WHERE id = 1").fetchone()
        return False if row is None else hmac.compare_digest(self._hash(password, row["salt"]), row["pw_hash"])

    def add(self, data: dict, photo: bytes = None):
        cols = FIELD_KEYS + ["photo"]
        vals = [data.get(k, "") for k in FIELD_KEYS] + [photo]
        sql = f"INSERT INTO records ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})"
        with self.conn:
            self.conn.execute(sql, vals)

    def update(self, rec_id, data: dict, photo: bytes = None):
        cols = FIELD_KEYS + ["photo"]
        vals = [data.get(k, "") for k in FIELD_KEYS] + [photo, rec_id]
        sql = f"UPDATE records SET {', '.join(c + '=?' for c in cols)} WHERE id=?"
        with self.conn:
            self.conn.execute(sql, vals)

    def delete(self, rec_id):
        with self.conn:
            self.conn.execute("DELETE FROM records WHERE id=?", (rec_id,))

    def get(self, rec_id):
        return self.conn.execute("SELECT * FROM records WHERE id=?", (rec_id,)).fetchone()

    def all(self):
        return self.conn.execute(
            "SELECT id, " + ", ".join(TABLE_COLS) + " FROM records ORDER BY id DESC").fetchall()

    def all_full(self):
        return self.conn.execute("SELECT * FROM records ORDER BY id ASC").fetchall()

    def search(self, keyword):
        like = f"%{keyword}%"
        where = " OR ".join(f"{k} LIKE ?" for k in FIELD_KEYS)
        return self.conn.execute(
            "SELECT id, " + ", ".join(TABLE_COLS) + f" FROM records WHERE {where} ORDER BY id DESC",
            [like] * len(FIELD_KEYS)).fetchall()


@st.cache_resource
def get_db() -> Database:
    return Database()


# ---------------------------------------------------------------- Helpers
def process_photo(file_like) -> bytes:
    img = Image.open(file_like)
    img = ImageOps.exif_transpose(img).convert("RGB")
    w, h = img.size
    target = PHOTO_SIZE[0] / PHOTO_SIZE[1]
    if w / h > target:
        new_w = int(h * target)
        left = (w - new_w) // 2
        img = img.crop((left, 0, left + new_w, h))
    else:
        new_h = int(w / target)
        top = (h - new_h) // 2
        img = img.crop((0, top, w, top + new_h))
    img = img.resize(PHOTO_SIZE, Image.LANCZOS)
    buf = BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def export_to_excel(rows, lang_code: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel Data"

    labels = [f[1][lang_code] for f in FIELDS]
    headers = ["No.", "Photo 3×4"] + labels
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=text)
        cell.font = Font(name="Phetsarath OT", bold=True)
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 36

    for n, row in enumerate(rows, start=1):
        r = n + 1
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value="")
        for j, key in enumerate(FIELD_KEYS, start=3):
            cell_val = ws.cell(row=r, column=j, value=row[key] or "")
            cell_val.font = Font(name="Phetsarath OT")
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center if c <= 2 else left
        ws.row_dimensions[r].height = 80
        if row["photo"]:
            xl_img = XLImage(BytesIO(row["photo"]))
            xl_img.width, xl_img.height = 75, 100
            ws.add_image(xl_img, f"B{r}")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    for c in range(3, len(headers) + 1):
        ws.column_dimensions[get_col(c)].width = 22
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def get_col(idx: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)


def valid_date(text: str) -> bool:
    if not text:
        return True
    if not re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", text):
        return False
    try:
        datetime.strptime(text, "%d/%m/%Y")
        return True
    except ValueError:
        return False


def t(key: str) -> str:
    return LANG[st.session_state.lang].get(key, key)


# ---------------------------------------------------------------- State
def init_state():
    defaults = {
        "lang": "lo",
        "authenticated": False,
        "failed_attempts": 0,
        "selected_id": None,
        "photo_bytes": None,
        "uploader_n": 0,
        "table_n": 0,
        "pending_load": None,      # dict ທີ່ຈະໂຫຼດເຂົ້າຟອມໃນຮອບຕໍ່ໄປ
        "confirm_delete": False,
        "flash": None,             # (type, message)
        "excel_bytes": None,
        "search_kw": "",
    }
    for k, v in defaults.items():
        st.session_state.setdefault(k, v)
    for k in FIELD_KEYS:
        st.session_state.setdefault(f"f_{k}", "")


def apply_pending_load():
    """ຕ້ອງເອີ້ນກ່ອນສ້າງ widget ເພື່ອໃຫ້ປ່ຽນຄ່າໃນຟອມໄດ້"""
    data = st.session_state.pending_load
    if data is None:
        return
    for k in FIELD_KEYS:
        st.session_state[f"f_{k}"] = data.get(k, "") or ""
    st.session_state.pending_load = None


def queue_clear():
    st.session_state.selected_id = None
    st.session_state.photo_bytes = None
    st.session_state.pending_load = {k: "" for k in FIELD_KEYS}
    st.session_state.uploader_n += 1
    st.session_state.table_n += 1
    st.session_state.confirm_delete = False


def flash(kind: str, msg: str):
    st.session_state.flash = (kind, msg)


# ---------------------------------------------------------------- Auth screens
def auth_screen(db: Database):
    _, mid, _ = st.columns([1, 2, 1])
    with mid:
        st.title(t("title"))
        st.selectbox(
            "🌐 " + t("language"), list(LANG_NAMES), index=list(LANG_NAMES).index(st.session_state.lang),
            format_func=LANG_NAMES.get, key="lang_auth",
            on_change=lambda: st.session_state.update(lang=st.session_state.lang_auth),
        )

        if not db.has_password():
            st.subheader(t("setup_title"))
            with st.form("setup"):
                pw1 = st.text_input(t("new_pw"), type="password")
                pw2 = st.text_input(t("confirm_pw"), type="password")
                if st.form_submit_button(t("save"), type="primary"):
                    if len(pw1) < MIN_PASSWORD_LEN:
                        st.error(t("pw_short"))
                    elif pw1 != pw2:
                        st.error(t("pw_mismatch"))
                    else:
                        db.set_password(pw1)
                        st.session_state.authenticated = True
                        st.rerun()
            return

        if st.session_state.failed_attempts >= MAX_LOGIN_ATTEMPTS:
            st.error(t("locked"))
            return

        st.subheader(t("login_title"))
        with st.form("login"):
            pw = st.text_input(t("password"), type="password")
            if st.form_submit_button(t("login"), type="primary"):
                if db.verify_password(pw):
                    st.session_state.authenticated = True
                    st.session_state.failed_attempts = 0
                    st.rerun()
                else:
                    st.session_state.failed_attempts += 1
                    left = MAX_LOGIN_ATTEMPTS - st.session_state.failed_attempts
                    st.error(f"{t('wrong_pw')} — {t('attempts')}: {left}")


# ---------------------------------------------------------------- Main app
def sidebar(db: Database):
    with st.sidebar:
        st.selectbox(
            "🌐 " + t("language"), list(LANG_NAMES), index=list(LANG_NAMES).index(st.session_state.lang),
            format_func=LANG_NAMES.get, key="lang_main",
            on_change=lambda: st.session_state.update(lang=st.session_state.lang_main),
        )
        st.divider()

        # Excel
        if st.button(t("export"), use_container_width=True):
            rows = db.all_full()
            if rows:
                st.session_state.excel_bytes = export_to_excel(rows, st.session_state.lang)
            else:
                st.session_state.excel_bytes = None
                st.info(t("no_data"))
        if st.session_state.excel_bytes:
            st.download_button(
                t("download"), st.session_state.excel_bytes,
                file_name=f"personnel_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        # Change password
        with st.expander(t("change_pw")):
            with st.form("chpw", clear_on_submit=True):
                old = st.text_input(t("old_pw"), type="password")
                n1 = st.text_input(t("new_pw"), type="password")
                n2 = st.text_input(t("confirm_pw"), type="password")
                if st.form_submit_button(t("save")):
                    if not db.verify_password(old):
                        st.error(t("wrong_pw"))
                    elif len(n1) < MIN_PASSWORD_LEN:
                        st.error(t("pw_short"))
                    elif n1 != n2:
                        st.error(t("pw_mismatch"))
                    else:
                        db.set_password(n1)
                        st.success(t("pw_updated"))

        st.divider()
        if st.button(t("logout"), use_container_width=True):
            for k in list(st.session_state.keys()):
                if k != "lang":
                    del st.session_state[k]
            st.rerun()


def read_form():
    data = {k: str(st.session_state.get(f"f_{k}", "") or "").strip() for k in FIELD_KEYS}
    if not data["name"]:
        st.warning(t("name_required"))
        return None
    for k in DATE_KEYS:
        if not valid_date(data[k]):
            label = next(f[1][st.session_state.lang] for f in FIELDS if f[0] == k)
            st.warning(f"{t('bad_date')} {label}")
            return None
    return data


def form_section(db: Database):
    lang = st.session_state.lang

    # ---- Photo (ເກັບຮູບທີ່ອັບໂຫຼດ)
    form_col, photo_col = st.columns([4, 1])

    with photo_col:
        st.markdown(f"**{t('photo')}**")
        up = st.file_uploader(
            t("upload_hint"), type=["jpg", "jpeg", "png", "bmp"],
            key=f"uploader_{st.session_state.uploader_n}", label_visibility="collapsed",
        )
        if up is not None:
            try:
                st.session_state.photo_bytes = process_photo(up)
            except Exception as e:
                st.error(f"{t('error')}: {e}")
        if st.session_state.photo_bytes:
            st.image(st.session_state.photo_bytes, use_container_width=True)
            if st.button(t("remove_photo"), use_container_width=True):
                st.session_state.photo_bytes = None
                st.session_state.uploader_n += 1
                st.rerun()
        else:
            st.caption(f"🖼 {t('no_photo')}")

    with form_col:
        title = t("form")
        if st.session_state.selected_id is not None:
            title += f"  —  {t('editing')}: {st.session_state.selected_id}"
        st.markdown(f"**{title}**")

        per_row = 3
        for start in range(0, len(FIELDS), per_row):
            cols = st.columns(per_row)
            for col, (key, labels, kind) in zip(cols, FIELDS[start:start + per_row]):
                label = labels[lang] + (" *" if key == "name" else "")
                with col:
                    if kind == "combo":
                        opts = [""] + COMBO_OPTIONS[key][lang]
                        cur = st.session_state.get(f"f_{key}", "")
                        if cur and cur not in opts:      # ຄ່າເກົ່າຈາກພາສາອື່ນ
                            opts.append(cur)
                        st.selectbox(label, opts, key=f"f_{key}")
                    elif kind == "date":
                        st.text_input(label, key=f"f_{key}", placeholder="DD/MM/YYYY")
                    else:
                        st.text_input(label, key=f"f_{key}")
        st.caption(t("required_hint"))

    # ---- Action buttons
    b1, b2, b3, b4, _ = st.columns([1, 1.4, 1, 1, 3])
    if b1.button(t("add"), type="primary", use_container_width=True):
        data = read_form()
        if data:
            db.add(data, st.session_state.photo_bytes)
            queue_clear()
            flash("success", t("added"))
            st.rerun()

    if b2.button(t("update"), use_container_width=True):
        if st.session_state.selected_id is None:
            st.info(t("select_row_err"))
        else:
            data = read_form()
            if data:
                db.update(st.session_state.selected_id, data, st.session_state.photo_bytes)
                queue_clear()
                flash("success", t("updated"))
                st.rerun()

    if b3.button(t("delete"), use_container_width=True):
        if st.session_state.selected_id is None:
            st.info(t("select_row_err"))
        else:
            st.session_state.confirm_delete = True

    if b4.button(t("clear"), use_container_width=True):
        queue_clear()
        st.rerun()

    if st.session_state.confirm_delete and st.session_state.selected_id is not None:
        st.warning(f"{t('confirm_delete')} (ID {st.session_state.selected_id})")
        c1, c2, _ = st.columns([1, 1, 6])
        if c1.button(t("yes"), type="primary"):
            db.delete(st.session_state.selected_id)
            queue_clear()
            flash("success", t("deleted"))
            st.rerun()
        if c2.button(t("cancel")):
            st.session_state.confirm_delete = False
            st.rerun()


def table_section(db: Database):
    lang = st.session_state.lang

    s1, s2, s3, _ = st.columns([3, 1, 1.3, 3])
    kw = s1.text_input(t("search"), value=st.session_state.search_kw,
                       placeholder=t("search_ph"), label_visibility="collapsed")
    if kw != st.session_state.search_kw:
        st.session_state.search_kw = kw
        st.session_state.table_n += 1
        st.rerun()
    s2.button(t("search"), use_container_width=True)  # ກົດເພື່ອ rerun
    if s3.button(t("show_all"), use_container_width=True):
        st.session_state.search_kw = ""
        st.session_state.table_n += 1
        st.rerun()

    kw = st.session_state.search_kw.strip()
    rows = db.search(kw) if kw else db.all()

    labels = {f[0]: f[1][lang] for f in FIELDS}
    df = pd.DataFrame(
        [[r["id"]] + [r[c] or "" for c in TABLE_COLS] for r in rows],
        columns=["ID"] + [labels[c] for c in TABLE_COLS],
    )
    ids = [r["id"] for r in rows]

    st.caption(f"{t('total')}: {len(rows)}")
    event = st.dataframe(
        df, hide_index=True, use_container_width=True, height=380,
        on_select="rerun", selection_mode="single-row",
        key=f"table_{st.session_state.table_n}",
    )

    sel = event.selection.rows if event and event.selection else []
    if sel:
        rec_id = ids[sel[0]]
        if rec_id != st.session_state.selected_id:
            rec = db.get(rec_id)
            if rec is not None:
                st.session_state.selected_id = rec_id
                st.session_state.photo_bytes = rec["photo"]
                st.session_state.pending_load = {k: rec[k] or "" for k in FIELD_KEYS}
                st.session_state.uploader_n += 1
                st.session_state.confirm_delete = False
                st.rerun()


def main_screen(db: Database):
    apply_pending_load()
    sidebar(db)
    st.title(t("title"))

    if st.session_state.flash:
        kind, msg = st.session_state.flash
        st.toast(msg, icon="✅" if kind == "success" else "⚠️")
        st.session_state.flash = None

    form_section(db)
    st.divider()
    table_section(db)


def main():
    st.set_page_config(page_title="Personnel DB", page_icon="🗂", layout="wide")
    init_state()
    db = get_db()
    if not st.session_state.authenticated:
        auth_screen(db)
    else:
        main_screen(db)


main()